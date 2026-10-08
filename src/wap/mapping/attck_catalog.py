from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from wap.models.schemas import TechniqueRef


def _is_yellow(cell, rgb_suffix: str = "FFFF00") -> bool:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return False
    rgb = getattr(fill.fgColor, "rgb", None)
    if rgb is None:
        return False
    return str(rgb).upper().endswith(rgb_suffix.upper())


def _parse_technique(cell_b: str) -> tuple[str, str]:
    """Split 'T1595  Active Scanning' -> (T1595, Active Scanning)."""
    s = (cell_b or "").strip()
    m = re.match(r"(T\d+(?:-[A-Za-z0-9]+)?)\s+(.*)$", s)
    if not m:
        raise ValueError(f"cannot parse technique cell: {cell_b!r}")
    return m.group(1), m.group(2).strip()


def load_attck_workbook(
    path: Path | str,
    sheet: str = "Sheet1",
    excluded_rgb_suffix: str = "FFFF00",
) -> list[TechniqueRef]:
    """
    Read filtered ATT&CK sheet.
    Yellow technique cell => retained=False (excluded from core coverage).
    Never modifies the workbook.
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"ATT&CK workbook not found: {p}")
    wb = openpyxl.load_workbook(p, data_only=True)
    if sheet not in wb.sheetnames:
        raise KeyError(f"sheet {sheet!r} missing; have {wb.sheetnames}")
    ws = wb[sheet]
    out: list[TechniqueRef] = []
    problems: list[str] = []
    seen: set[str] = set()
    for r in range(2, ws.max_row + 1):
        tactic = ws.cell(r, 1).value
        tech = ws.cell(r, 2).value
        meaning = ws.cell(r, 3).value or ""
        if not tech:
            continue
        try:
            tid, tname = _parse_technique(str(tech))
        except ValueError as e:
            problems.append(str(e))
            continue
        # same technique may appear under multiple tactics — keep each (tactic,id) once
        key = f"{tid}::{tactic}"
        if key in seen:
            problems.append(f"duplicate row for {key}")
            continue
        seen.add(key)
        # note: technique_id alone is NOT unique across tactics
        excluded = _is_yellow(ws.cell(r, 2)) or _is_yellow(ws.cell(r, 1))
        out.append(
            TechniqueRef(
                technique_id=tid,
                technique_name=tname,
                tactic=str(tactic or ""),
                description=str(meaning),
                retained=not excluded,
            )
        )
    if problems:
        # surface but do not silent-fix
        print("[attck_catalog] warnings:")
        for pmsg in problems[:30]:
            print(" -", pmsg)
        if len(problems) > 30:
            print(f" - ... and {len(problems)-30} more")
    return out


def catalog_summary(techs: list[TechniqueRef]) -> dict:
    retained = [t for t in techs if t.retained]
    excluded = [t for t in techs if not t.retained]
    return {
        "total_rows": len(techs),
        "retained": len(retained),
        "excluded": len(excluded),
        "unique_retained_ids": len({t.technique_id for t in retained}),
        "unique_excluded_ids": len({t.technique_id for t in excluded}),
    }